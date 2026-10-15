from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ASSETS = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[2] / 'docs/status_strip'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ('Heading 1', 'Heading 2'):
    doc.styles[name].paragraph_format.space_before = Pt(8)
    doc.styles[name].paragraph_format.space_after = Pt(6)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Grinder 控制盒状态灯带  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)

def p(text, style=None):
    return doc.add_paragraph(text, style)

pending_page = False

def h(text):
    global pending_page
    heading = doc.add_heading(text, 1)
    if pending_page:
        heading.paragraph_format.page_break_before = True
        pending_page = False

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width); cell.text = label
        for run in cell.paragraphs[0].runs: run.bold = True
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9EEF2'); cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width); cell.text = str(text)
    for row in t.rows:
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for cell in row.cells:
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + side); node.set(qn('w:w'), '80'); node.set(qn('w:type'), 'dxa'); margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs: run.font.size = Pt(10)
    spacer = p('')
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    return t

def code(lines):
    for line in lines.splitlines():
        para = p(line)
        para.paragraph_format.space_after = Pt(1)
        for run in para.runs:
            run.font.name = 'Consolas'; run.font.size = Pt(9)

def page():
    global pending_page
    pending_page = True

doc.add_heading('RK3588 控制盒状态灯带实施方案', 0)
p('适用硬件  ESP32-S3-Zero-N8R8-M 与三颗 WS2812B 灯珠')
p('版本 1.2    日期 2026年10月14日')
p('本方案用于给 Grinder 控制盒增加状态指示。采用 RK3588 USB 向 ESP32 提供电源和通信，ESP32 通过 GPIO4 控制 10cm 灯带。选购成品 5V WS2812B RGB 灯带，30 灯每米，截取 3 颗灯；软件亮度限制为 20%。灯带只提示状态，不承担急停或停车保护。')
h('一 采购清单')
table(['部件', '具体规格与数量'], [
('控制板', '微雪 ESP32-S3-Zero-N8R8-M，SKU 33549，已焊排针版，1 块。8MB Flash 与 8MB PSRAM。'),
('灯带', '5V WS2812B RGB，5050 灯珠，30 灯/米，单线 800kHz，三线 5V／DIN／GND；买成品带并截取 10cm、3 颗灯。建议 IP30 裸板置于保护扩散罩内。'),
('电平转换', 'TI SN74AHCT125DR，SOIC-14 贴片版，1 颗，立创编号 C155176。按本次选择直接焊细线，不使用直插版或转接板。'),
('数据电阻', '330Ω、1/4W（250mW）、±5%，1 只。截图型号 CR1/4W-330Ω±5%-OT52，编号 C2894590，规格适用；留意该商品的起售数量。'),
('连接附件', '支持数据传输的 USB Type-C 线 1 根；短细导线、热缩管、绝缘固定材料与乳白扩散罩。'),
], [1.15, 5.85])
p('按本次要求不装去耦电容或入口滤波电容。这是不加电容的试装配置，尚未验证稳定性；USB 来自 ESP32 并不等同于芯片就近去耦，TI 仍建议在 VCC/GND 附近安装 100nF。[2]')
p('给卖家的采购描述：5V WS2812B RGB 幻彩灯带，30 灯每米，5050，三线单信号，GRB 顺序，单颗可裁剪，带输入端引线。不要买 12V 灯带、四线模拟 RGB 灯带或 RGBW 灯带。')
p('10cm 只有 3 颗灯，适合局部状态窗。若要围满外壳一整圈，先量周长，再重新确定灯数、扩散效果和供电能力；当前固件固定 3 颗，不能仅加长灯带就认为所有灯都会亮。')

page()
h('二 GPIO4 的实际位置')
p('板子元件正面朝向你，Type-C 接口朝上。GPIO4 在左侧排针，认准板边丝印数字 4，位于 3 与 5 之间。GPIO 编号不是从上往下数的第 4 根针，也不是芯片本体的第 4 脚。')
im = Image.open(ASSETS / 'pinout.webp').convert('RGB'); im.save(ASSETS / 'pinout.png')
doc.add_picture(str(ASSETS / 'pinout.png'), width=Inches(6.0))
p('图 1  微雪官方 ESP32-S3-Zero 系列引脚图。上半部为正面；下半部为背面。接线以实物丝印和对应型号官方资料为准。[1]')
table(['板上丝印', '本方案用途'], [
('4', 'GPIO4 数据输出 → 74AHCT125 的 2 脚 1A。'),
('5V', 'USB 的 5V 电源支路 → 灯带正极及 74AHCT125 的 14 脚。'),
('GND', '公共地 → 灯带负极、74AHCT125 的 7 脚与 1 脚。'),
('3V3', '本方案不接灯带电源。GPIO4 本身输出的是 3.3V 逻辑信号。'),
], [1.15, 5.85])
p('板载小灯使用 GPIO21；它和外接灯带的 GPIO4 是两个不同信号。不要把 TX、RX、USB D+／D− 当作灯带数据口。')

page()
h('三 接线与供电')
p('先断开 USB 后接线。USB 是唯一电源来源；5V 电源经板上 5V 引脚分给灯带与转换芯片。灯带电流不经过 GPIO，也不经过板上的 3.3V 稳压输出。不要同时向板子并接另一组外部 5V。')
code('RK3588 USB ── Type-C 数据线 ── ESP32-S3-Zero\nESP32 5V ──────┬── 灯带 5V\n               └── 74AHCT125 14脚 VCC\nESP32 GND ─────┬── 灯带 GND\n               ├── 74AHCT125 7脚 GND\n               └── 74AHCT125 1脚 /1OE\nESP32 GPIO4 ────── 74AHCT125 2脚 1A\n74AHCT125 3脚 1Y ── 330Ω ── 灯带 DIN → 第1颗 → 第2颗 → 第3颗')
p('灯带箭头指向数据流动方向。接 DIN 输入端，DOUT 末端留空；以焊盘文字为准，不依赖红黄黑线色判断。330Ω 靠近第一颗灯的 DIN 焊盘安装；它用于抑制数据尖峰，不串在 5V 电源线上。')
doc.add_heading('330Ω 电阻接法', 2)
p('电阻只有两根引线，没有正负极，两端可互换。一根引线焊到 SN74AHCT125DR 的 3 脚 1Y（可用短导线连接），另一根引线焊到第一颗灯的 DIN 输入焊盘。串联后用热缩管绝缘，电阻尽量靠近 DIN。它不接 5V，不接 GND，也不跨接在 5V 与 GND 之间。')
code('芯片 3脚 1Y ── 电阻任意一端 [330Ω] 另一端 ── 第一颗灯 DIN')
table(['74AHCT125 引脚', '连接位置'], [
('14  VCC', 'ESP32 的 5V。'),
('7  GND 与 1  /1OE', '接公共 GND；/1OE 低电平使能第一通道。'),
('2  1A 与 3  1Y', '2 脚接 GPIO4；3 脚经 330Ω 接灯带 DIN。'),
('4、10、13  未用 /OE', '接 5V，关闭未用通道。'),
('5、9、12  未用 A', '接 GND，避免输入悬空。'),
('6、8、11  未用 Y', '不连接。'),
], [2.0, 5.0])
p('本方案使用 SOIC-14 贴片封装。从芯片顶面看，缺口朝上或 1 脚标记位于左上时，左边自上而下为 1～7 脚，右边自下而上为 8～14 脚，按实物标记确认。[2] 3 脚就是左侧从 1 脚向下的第三脚，不是 ESP32 的 GPIO3。芯片 5V 来自同一 USB，不需额外电源。焊细线时防止相邻脚短接，并做绝缘和机械固定。')
p('导线尽量短，建议控制器到第一颗灯不超过 20cm，并让数据线靠近地线。按传统每颗满白约 60mA 的保守值估算，3 颗灯满亮约 180mA；20% 满白约 36mA，仅为灯珠估算，未含 ESP32 和静态损耗。实际批次、USB 端口及线材需上电测量，不能把开发板 3.3V 稳压器标称电流当成 USB 余量。[3]')

page()
h('四 固件烧录与板端启用')
p('先仅连接 ESP32 与灯带验证，再接入 RK3588。仓库固件位于 firmware/status_strip，驱动引脚固定 GPIO4，灯数固定 3，GRB／800kHz，亮度上限 51/255。板端不需要安装 PlatformIO。')
doc.add_heading('N8R8 的构建设置', 2)
p('现有 platformio.ini 按基础版 4MB Flash 设置。N8R8 具有 8MB Flash；建议将下面三项改为匹配该型号的配置，再重新构建；保留项目现有其余配置，不启用 PSRAM。该版本尚未在实物上烧录验收。')
code('board_upload.flash_size = 8MB\nboard_upload.maximum_size = 8388608\nboard_build.partitions = default_8MB.csv')
p('开发电脑在 firmware/status_strip 目录执行，COM 端口必须替换为 ESP32 实际端口：')
code('pio run\npio run -t upload --upload-port COM实际编号')
p('首次下载可按住 BOOT 后插入 USB，必要时按 RESET，再释放 BOOT。使用可传数据的 USB 线。烧录完成重启后，应看到 3 颗外接灯黄灯常亮。不要把底盘串口当作烧录端口。')
doc.add_heading('接入 RK3588', 2)
p('插入 ESP32 后，在板端查询 USB 身份：')
code('python3 -m serial.tools.list_ports -v\nsudo nano /etc/grinder/startup.conf')
p('保留已经验证的主系统路径与用户配置，将以下字段填好。USB 序列号与稳定设备路径二选一；示例不是真实序列号：')
code("START_INDICATOR=1\nINDICATOR_USB_SERIAL_NUMBER='实际SER值'\nINDICATOR_DEVICE=''")
p('确认任务停止，再逐条执行：')
code('sudo grinder-service apply\nsudo grinder-service restart\ngrinder-service status')
p('已有服务会统一启动灯带，不要另开一个 status_indicator.launch 抢占 USB。服务使用 dialout 补充组。没有 ESP32 时仍保持 START_INDICATOR=0；开机自启开关 AUTOSTART 与灯带开关相互独立。')
p('日常命令：sudo grinder-service enable 开启下次自启；disable 关闭下次自启；restart 重启节点；sudo reboot 重启板端。更完整的软件部署步骤见仓库根 README。')

page()
h('五 灯效与上电验收')
table(['运行状态', '三颗灯同步显示'], [
('通电但 ROS 状态尚不完整', '黄色常亮'),
('待机 准备就绪 正常停止 路径规划', '绿色常亮'),
('自动任务运行', '蓝色呼吸，周期 2 秒'),
('人为暂停', '蓝色慢闪，1Hz'),
('RPP 碰撞风险导致任务暂停', '黄色慢闪，1Hz；风险消失不自动清除此暂停原因'),
('手动活动或非自动任务磨盘使能', '青色常亮；不能据此认定已测得实际转速'),
('任务完成', '绿色闪三次，然后常亮'),
('故障 底盘断连 就绪后状态过期', '红色慢闪，1Hz'),
('已收到且未清除的急停提示', '红色快闪，4Hz，优先于其他灯效'),
('USB 断电', '熄灭；若仍供电但已建立过通信后中断超过 3 秒，转红色慢闪'),
], [3.65, 3.35])
p('验收顺序：① 断电检查 5V/GND 无短路、贴片芯片脚间无焊锡桥，电阻串接在 3 脚与 DIN 之间；② 固件烧录后单独上电，3 颗均黄灯；③ 接入已启动的主系统，ROS 状态完整后绿灯；④ 检查蓝色任务灯与两种暂停颜色；⑤ 停止通信但保留电源，验证超时红灯；⑥ 确认任务停止后重启板端，检查 enabled 与 active (running)。涉及轮速、磨盘和急停的验证需按设备现场操作流程完成。')
p('出现不亮时先测灯带输入 5V 和共地，再查 DIN 方向、GPIO4 丝印、/1OE 接地及固件灯数；有错色时核对 GRB 顺序。仅首颗亮或闪烁时检查焊点、供电、线长和电平转换。三颗灯经扩散罩仍可见亮斑，不保证形成连续均匀的一整圈光带。')
doc.add_heading('资料与验证范围', 2)
for text in [
'[1] 微雪 ESP32-S3-Zero 系列产品说明及引脚图\nhttps://docs.waveshare.net/ESP32-S3-Zero/',
'[2] TI SN74AHCT125 官方数据手册\nhttps://www.ti.com/lit/ds/symlink/sn74ahct125.pdf',
'[3] Adafruit NeoPixel 使用建议与供电说明\nhttps://learn.adafruit.com/adafruit-neopixel-uberguide/best-practices\nhttps://learn.adafruit.com/adafruit-neopixel-uberguide/powering-neopixels',
]:
    para=p(text)
    for run in para.runs: run.font.size=Pt(8.5)
p('资料核对日期 2026年10月14日。RK3588 主系统开机自启已通过用户现场验证；ESP32-N8R8、灯带供电、电平波形和真实灯效尚待硬件验收。')
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(ROOT / 'RK3588控制盒灯带接线与部署方案.docx')
print(ROOT / 'RK3588控制盒灯带接线与部署方案.docx')
