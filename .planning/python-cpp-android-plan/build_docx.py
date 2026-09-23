from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\CODE\C++\Grinder")
OUT = ROOT / "artifacts" / "Grinder设备端与安卓端Python与C++加速实施方案.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2F5597"
PALE = "D9EAF7"
PALE2 = "EEF4F8"
GRAY = "F2F2F2"
MIDGRAY = "D9E1F2"
TEXT = RGBColor(31, 41, 55)
RED = RGBColor(166, 31, 31)
GREEN = RGBColor(38, 103, 71)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="C9D2DC", size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    paragraph.add_run(" 页")


def add_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
    for ri, row in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            if ri % 2:
                set_cell_shading(cell, "F7F9FB")
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(text))
            r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(2)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_callout(doc, title, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="9DB2C8", size="7")
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow(doc, labels):
    table = doc.add_table(rows=1, cols=len(labels) * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, label in enumerate(labels):
        c = table.cell(0, i * 2)
        c.width = Inches(1.02)
        set_cell_shading(c, BLUE if i in (0, len(labels) - 1) else PALE)
        set_cell_border(c, color="8FAADC")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(7.2)
        if i in (0, len(labels) - 1):
            r.font.color.rgb = RGBColor(255, 255, 255)
        if i < len(labels) - 1:
            a = table.cell(0, i * 2 + 1)
            a.width = Inches(0.27)
            p2 = a.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = p2.add_run("→")
            rr.bold = True
            rr.font.size = Pt(13)
            set_cell_border(a, color="FFFFFF", size="0")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.88)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)
sec.header_distance = Inches(0.25)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(9.3)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(4)
styles["Normal"].paragraph_format.line_spacing = 1.14
for name, size, color in [("Title", 26, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 10.5, NAVY)]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

for s in doc.sections:
    h = s.header.paragraphs[0]
    h.text = "GRINDER 设备端 × GrindingRobot APP｜工程实施方案"
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs:
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(108, 117, 125)
    add_page_number(s.footer.paragraphs[0])
    for r in s.footer.paragraphs[0].runs:
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(108, 117, 125)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(55)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Grinder 设备端与安卓端")
r.bold = True
r.font.size = Pt(27)
r.font.color.rgb = RGBColor.from_string(NAVY)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Python 保持路线与 C++ 热路径加速实施方案")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor.from_string(BLUE)

doc.add_paragraph()
add_callout(
    doc,
    "核心结论",
    "原方案具有较强参考价值，但不能原样执行。当前代码已完成设备端 32KB 接收、4KB 地图分片和部分有界队列；最高优先级应转为 Android 单写者、地图传输事务化、Super-LIO 健康门槛，以及安全与可观测性。C++ 只在基准确认热点后进入。",
)
meta = add_table(
    doc,
    ["项目", "内容"],
    [
        ["适用范围", "Grinder20260901 / ROS1 Noetic 设备端；GrindingRobot Android APP"],
        ["评审日期", "2026-09-21"],
        ["输出性质", "代码现状核对 + 分阶段工程实施 + 测试与验收方案"],
        ["推荐路线", "先 Python/Kotlin/协议治理，再按火焰图结果选择性引入 C++"],
        ["不在本次范围", "不直接改动生产业务代码；不替换 Super-LIO、MST27 或完整调度器"],
    ],
    widths=[1.35, 5.85],
    font_size=8.6,
)
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("内部工程评审稿｜建议基线冻结后进入实施")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(100, 110, 120)
doc.add_page_break()

doc.add_heading("文档导读", level=1)
add_table(
    doc,
    ["章节", "回答的问题"],
    [
        ["1. 综合判断", "原方案哪些正确、哪些已完成、哪些需要纠偏？"],
        ["2. 当前代码事实", "设备端和 Android 现在真正的瓶颈在哪里？"],
        ["3. 目标架构", "控制、事务与大数据如何隔离？"],
        ["4. 路线一", "不改 C++ 时具体改什么、先后顺序是什么？"],
        ["5. 路线二", "何时值得 C++ 化，边界和接口如何设计？"],
        ["6. 实施计划", "6 周如何落地、如何灰度和回滚？"],
        ["7. 测试验收", "功能、性能、故障、安全与稳定性如何证明？"],
        ["8. 风险与交付物", "上线门禁、责任边界和最终产物是什么？"],
    ],
    widths=[1.3, 5.9],
)
add_callout(doc, "阅读建议", "管理评审重点看第 1、6、7 节；设备端开发重点看第 2、4、5 节；Android 开发重点看第 2.3、4.3 和 7.2 节。", fill=PALE2)

doc.add_heading("1. 综合判断：有参考价值，但需按现状重排优先级", level=1)
p = doc.add_paragraph()
p.add_run("总体评价：").bold = True
p.add_run("约 70% 的方向成立，特别是“先测量、协议兼容优先、控制与大数据隔离、事务使用 operation_id、不要整体翻译 scheduler”这些原则。问题在于若干现状判断来自旧版本，且缺少 Android 发送并发、地图组包资源边界、安全门禁与定位质量门槛。")

add_table(
    doc,
    ["原方案主张", "当前事实", "判断", "修订建议"],
    [
        ["设备 TCP 接收缓冲 4096B", "设备端已使用 recv(32KB)", "已完成/过时", "不再作为设备端任务；Android 仍为 4096B，可调至 16–32KB"],
        ["地图分片被压到 512B", "设备端和 APP 当前均使用 4KB 上限", "已完成/过时", "先优化发送和组包，再 A/B 测试 16KB"],
        ["批量发送地图", "路径响应已有约 128KiB 合并写；地图仍逐片同步发送", "部分成立", "地图改为流式窗口 + 批量写，禁止一次构造所有分片"],
        ["生产关闭逐帧详细 RX 日志", "设备端已偏向头部日志；Android 发送成功仍构造全帧十六进制字符串", "仍需完成", "双端统一采样日志和统计日志"],
        ["单一发送线程/队列", "设备端有 send_lock；Android 普通发送、异步发送和心跳可并发 write+flush", "高优先级", "Android 先建单写者优先级 Channel"],
        ["Livox 重复启动", "当前 mapping launch 引入 Super-LIO，不是 Livox 驱动；base 脚本启动驱动", "当前不成立", "保留唯一所有权与节点冲突检查，避免回归"],
        ["C++ 加速 SL 帧和 Protobuf", "可能有效，但尚无火焰图证明；Python Protobuf/OpenCV已有原生实现", "有条件成立", "达到热点阈值后再做，不先上 JNI/全量重写"],
    ],
    widths=[1.45, 2.0, 1.05, 2.7],
    font_size=7.7,
)

add_callout(
    doc,
    "决策",
    "路线一必须先实施。路线二不是路线一失败后的补救，而是通过同一套基准证明某个 CPU 热点仍无法达标时，对该热点进行可回滚替换。",
)

doc.add_heading("2. 当前代码事实与瓶颈", level=1)
doc.add_heading("2.1 端到端链路", level=2)
add_flow(doc, ["Android APP", "SL-LinkA / TCP", "grinder_scheduler", "ROS / Super-LIO", "底盘与地图"])
p = doc.add_paragraph()
p.add_run("性能问题必须拆成五段：").bold = True
p.add_run(" APP 排队与组包、TCP 传输、设备端解析与排队、ROS 业务执行、结果持久化/反馈。只看设备总 CPU 无法判断是否需要 C++。")

doc.add_heading("2.2 设备端已具备的能力", level=2)
add_bullets(doc, [
    "SL-LinkA 接收使用 32KB 缓冲，存在 send_lock、有界工作队列、控制响应并发限制与批量响应并发限制。",
    "地图 MapRequest 分片默认及上限为 4096B，外层帧仍受 uint16 长度约束，4KB处于安全区间。",
    "路径批量响应已能把相邻完整协议帧合并为约 128KiB 的 TCP 写入，同时保持每帧可独立解析。",
    "Super-LIO 模式管理已有串行操作队列、启动服务/话题等待、残留建图节点 kill 兜底和地图文件基本校验。",
])

doc.add_heading("2.3 仍然存在的主要问题", level=2)
add_table(
    doc,
    ["优先级", "问题", "代码证据/行为", "主要影响"],
    [
        ["P0", "Android 并发写同一 OutputStream", "sendData、sendDataAsync、heartbeat 分别 launch 后 write+flush", "帧交错、发送顺序不确定、频繁 flush、错误状态难追踪"],
        ["P0", "地图传输缺少事务标识和校验", "MapChunk 只有 map_id/map_version/index/count，无 transfer_id、total_bytes、整图摘要", "并发请求串图、旧分片污染、完整性不可证明"],
        ["P0", "生命周期和定位 READY 门槛不足", "停止 kill 后不复核；initialpose 后任意 /lio/odom 即 ready", "APP 误报成功、残留节点、低质量定位进入导航"],
        ["P0", "安全与鉴权基线不足", "控制服务暴露、生产凭据/明文链路风险、底盘反馈闭环不足", "未授权控制和故障状态下运动风险"],
        ["P1", "设备地图按请求重复全量处理", "OccupancyGrid→NumPy→RGB→flip/resize→PNG；每次重做", "大图首包延迟与CPU抖动"],
        ["P1", "设备地图先构造全部分片", "outputs 列表持有全部序列化分片；随后逐片同步发送", "内存峰值、接收线程占用、控制消息尾延迟"],
        ["P1", "Android 地图全内存重组", "TreeMap保存所有ByteArray，完成后再写入ByteArrayOutputStream并复制", "大地图内存峰值、OOM/GC风险"],
        ["P1", "帧解析存在额外复制", "Android CRC重新拼接header+payload，完成帧再copy payload", "高吞吐下对象分配和GC"],
        ["P1", "测试覆盖缺口", "缺TCP单写者、地图缺片/乱序/超时/恶意长度、断线重连压力测试", "优化后容易出现隐性兼容回归"],
    ],
    widths=[0.55, 1.45, 3.2, 2.0],
    font_size=7.5,
)

doc.add_heading("2.4 与通信性能同等重要的系统风险", level=2)
add_bullets(doc, [
    "三维建图内存并非主要由 Python 造成：Super-LIO 存在累计点云及关键帧全量复制路径，地图规模扩大时会直接推高 RSS。",
    "定位链路还需四元数归一化、空目标点云防护、ICP/NDT fitness 与连续稳定帧门槛，否则“响应快”不等于“可以导航”。",
    "MST27 已有 C++ 热核；规划更需要超时、取消、内存预算和大图降采样，而不是重复翻译算法。",
    "任何通信优化不得绕过急停、task_enable、速度清零、底盘反馈和控制仲裁。",
])

doc.add_heading("3. 目标架构", level=1)
doc.add_heading("3.1 三类业务、三套语义", level=2)
add_table(
    doc,
    ["通道", "典型消息", "队列策略", "确认语义", "优先级"],
    [
        ["实时控制", "急停、摇杆、速度、暂停", "有界；速度 latest-wins；急停不可丢", "接收/应用确认；超短超时", "最高"],
        ["业务事务", "建图、保存、定位、删除、模式切换", "可靠 FIFO；幂等键；禁止忙时静默丢弃", "accepted→running→terminal", "中"],
        ["大数据", "地图、轨迹、图片、历史记录", "窗口化流式传输；可取消；限内存", "分片进度+完整性校验", "低"],
    ],
    widths=[1.0, 1.6, 2.25, 1.7, 0.65],
)

doc.add_heading("3.2 单连接发送模型", level=2)
add_flow(doc, ["业务生产者", "优先级有界队列", "单写者", "批量聚合", "TCP"])
add_bullets(doc, [
    "Android 和设备端都只允许一个逻辑写者拥有 socket 输出流；其他协程/线程只提交不可变帧。",
    "急停/控制可抢占尚未出队的大数据，但不能把一个已经开始写入的协议帧切开。",
    "按字节数而不是仅按消息数限制队列，防止大地图分片耗尽内存。",
    "flush 由单写者按批次或延迟阈值执行；心跳也作为普通高优先级帧进入队列。",
])

doc.add_heading("3.3 地图传输状态机", level=2)
add_flow(doc, ["REQUEST", "ACCEPTED", "STREAMING", "VERIFYING", "COMPLETED / FAILED"])
doc.add_page_break()
add_table(
    doc,
    ["字段", "用途", "兼容策略"],
    [
        ["transfer_id", "唯一标识一次下载，避免同 map_id 并发污染", "新增字段；旧端为空时退化为 map_id+version"],
        ["map_id + map_version", "标识业务对象及不可变版本", "保留现有字段"],
        ["chunk_index / total_chunks", "乱序重组与进度", "保留并增加范围校验"],
        ["total_bytes", "提前设置资源上限，识别截断", "新增可选字段"],
        ["sha256 或 CRC32", "整图完成后校验", "优先SHA-256；低端兼容CRC32"],
        ["encoding / geometry", "PNG、栅格和坐标元数据", "首片或manifest固定，后续不得变化"],
        ["cancel / resume_token", "取消和可选断点续传", "先实现取消，断点续传后置"],
    ],
    widths=[1.35, 3.5, 2.35],
    font_size=8.0,
)

doc.add_heading("4. 路线一：Python、Kotlin 与协议优化（推荐先做）", level=1)
doc.add_heading("4.1 基线与可观测性", level=2)
add_numbered(doc, [
    "冻结测试场景：同一设备、同一雷达频率、同一地图文件、同一 APP 构建、同一日志级别。",
    "为每次请求生成 request_id/operation_id/transfer_id，记录 monotonic 时间戳。",
    "分别统计 APP 入队、APP 写完成、设备收帧、业务开始/结束、首片/末片发送、APP 校验/显示。",
    "每秒输出聚合指标：队列字节数、帧率、丢弃数、发送批次、p50/p95/p99、RSS、GC、CPU。",
    "采集设备 perf/py-spy、Android Perfetto/Studio Profiler 和网络 pcap，形成可复现基线报告。",
])

doc.add_heading("4.2 设备端实施项", level=2)
add_table(
    doc,
    ["序号", "改造项", "实现要点", "预期收益", "风险"],
    [
        ["D1", "地图响应接入 bulk TX", "复用现有批量写；控制队列优先；队列按字节限额", "减少 sendall/锁竞争，降低控制尾延迟", "低"],
        ["D2", "分片生成器/窗口", "不返回完整 outputs；按窗口生成、序列化、发送、释放", "内存不随分片总数线性增长", "中"],
        ["D3", "版本化地图缓存", "key=map_id+version+encoding+尺寸/旋转参数；保存时生成；原子替换", "消除重复PNG编码，缩短首包", "中"],
        ["D4", "日志降频", "生产默认只记录错误和每秒统计；禁止payload JSON化", "降低CPU与磁盘抖动", "低"],
        ["D5", "事务状态机", "统一operation_id和终态；超时不得返回success", "消除APP假成功", "中"],
        ["D6", "启动健康门槛", "PID+ROS节点+服务+话题频率+TF+定位质量", "降低拉不起/关不掉和假READY", "中"],
    ],
    widths=[0.45, 1.25, 2.9, 1.9, 0.7],
    font_size=7.5,
)

doc.add_page_break()
doc.add_heading("4.3 Android 实施项", level=2)
add_table(
    doc,
    ["序号", "改造项", "实现要点", "预期收益", "风险"],
    [
        ["A1", "单写者发送器", "CoroutineScope+Channel；控制/事务/大数据优先级；所有write/flush唯一入口", "先解决正确性，再降低系统调用", "中"],
        ["A2", "接收缓冲 16–32KB", "保留任意分片/粘包兼容；复用缓冲或ByteBuffer", "减少read与copyOf次数", "低"],
        ["A3", "解析器少复制", "增量CRC；复用payload buffer；限制最大帧/解析预算", "降低分配与GC", "中"],
        ["A4", "地图流式组包", "按transfer_id写临时文件；位图/BitSet记录分片；校验后原子发布", "内存峰值稳定，支持大图", "中"],
        ["A5", "资源与超时限制", "total_bytes/chunks上限、空闲超时、重复片幂等、并发传输数限制", "防OOM与恶意/异常输入", "低"],
        ["A6", "日志治理", "禁止生产全帧hex；采样+计数；敏感字段脱敏", "明显减少UI/IO抖动", "低"],
    ],
    widths=[0.45, 1.3, 3.0, 1.8, 0.65],
    font_size=7.5,
)

doc.add_heading("4.4 Super-LIO 生命周期与定位质量", level=2)
add_bullets(doc, [
    "停止：shutdown→等待→仅清理自有节点→再次确认；仍有残留则进入 TIMEOUT/ERROR 并返回残留节点清单。",
    "启动：服务存在只是 STARTING；关键话题需满足频率与时间戳新鲜度，TF需可变换且四元数有效。",
    "定位：initialpose 必须与当前 map_id/version 关联；连续 N 帧配准质量、fitness、协方差和位姿跳变均达标后才 READY。",
    "保存：PCD、PGM、YAML、元数据写入唯一临时目录，fsync 后原子切换；保存成功与定位启动成功分开反馈。",
    "驱动所有权：保留 Livox 唯一启动入口；mapping launch 只拥有 Super-LIO/mapper/loop，启动前检查同名节点与端口。",
])

doc.add_heading("5. 路线二：仅对高频热路径做 C++ 加速", level=1)
doc.add_heading("5.1 启动条件", level=2)
add_callout(
    doc,
    "进入 C++ 路线的门槛",
    "完成路线一后，在三轮可重复压测中，SL帧扫描/CRC/组帧/Protobuf编解码合计仍占设备单核 10%–15% 以上，或控制 p99/地图吞吐仍未达标且瓶颈明确位于 CPU，而不是网络、磁盘、PNG编码或ROS业务，才进入C++实现。",
)

doc.add_heading("5.2 推荐边界", level=2)
add_table(
    doc,
    ["模块", "建议", "边界/接口", "原因"],
    [
        ["SL 帧扫描、CRC、pack/unpack", "优先候选", "C++库/pybind11扩展；输入bytes，输出帧头+payload view", "纯计算、高频、边界清晰"],
        ["Protobuf 编解码", "条件候选", "只覆盖高频固定消息；协议生成物同源", "Python Protobuf本身可能已足够快"],
        ["批量发送/地图分片", "先Python实现", "memoryview+生成器+sendall批量窗口", "主要收益来自模型改变，不是语言"],
        ["Android 解析器 JNI", "暂不建议", "仅在Perfetto证明解析持续热点后考虑", "JNI增加ABI、崩溃和生命周期复杂度"],
        ["scheduler_node.py", "不整体重写", "业务状态机保留Python", "规则复杂、变更频繁，翻译收益低"],
        ["super_lio_mode_manager.py", "不整体重写", "ROS编排保留Python，补状态机和健康检查", "瓶颈是语义和可靠性"],
        ["MST27 / Super-LIO", "继续现有C++", "补预算、健康门槛和内存上限", "核心算法已经是C++路径"],
    ],
    widths=[1.55, 1.1, 2.35, 2.2],
    font_size=7.7,
)

doc.add_heading("5.3 C++ 组件工程要求", level=2)
add_bullets(doc, [
    "协议源唯一：只修改 sl_link.proto，再统一生成 Python、Android 与 C++ 代码；禁止手改生成文件。",
    "兼容性：旧 Python codec 与新 C++ codec 进行字节级 golden corpus 对比，包含空载荷、最大载荷、未知字段和错误CRC。",
    "内存：固定上限、RAII、无未界定所有权；禁止把 payload 指针跨越输入缓冲生命周期。",
    "故障隔离：扩展加载失败可通过 feature flag 回退 Python；C++异常不得越过Python边界。",
    "发布：x86_64 与 aarch64 双平台 Release 构建，固定编译器/ABI，输出符号和版本信息。",
])

doc.add_heading("6. 分阶段实施计划", level=1)
add_table(
    doc,
    ["阶段", "时间", "设备端", "Android", "出口条件"],
    [
        ["0 基线冻结", "3–5天", "埋点、perf/py-spy、网络抓包", "Perfetto、GC/内存、端到端时间戳", "可重复基线误差≤10%"],
        ["1 正确性优先", "1周", "事务分类、地图走bulk、停止复核", "单写者、生产日志降频", "10万帧无交错；急停优先"],
        ["2 地图传输", "1–1.5周", "缓存、流式窗口、transfer manifest", "流式落盘、校验、超时/取消", "2/10/50MB全通过故障注入"],
        ["3 生命周期", "1周", "operation状态机、健康门槛、定位质量", "分阶段状态UI、重试/取消", "无假成功、残留可诊断"],
        ["4 性能优化", "1周", "4KB/16KB A/B、复制优化", "16/32KB接收、解析少复制", "目标SLO达标且RSS稳定"],
        ["5 C++决策", "2–3天", "审查火焰图与CPU占比", "评估是否无需JNI", "满足门槛才立项"],
        ["6 灰度上线", "1周", "feature flag、10%→50%→100%", "兼容旧协议与回滚", "24h稳定+现场回归"],
    ],
    widths=[1.05, 0.75, 2.25, 2.15, 1.3],
    font_size=7.4,
)

doc.add_heading("6.1 推荐任务拆分", level=2)
add_numbered(doc, [
    "提交 1：只加指标和测试夹具，不改行为；形成基线报告。",
    "提交 2：Android 单写者和日志治理；保留旧实现 feature flag。",
    "提交 3：设备地图响应接入 bulk + 生成器窗口；协议不变。",
    "提交 4：协议新增 transfer_id/total_bytes/checksum/cancel；双端生成物同步。",
    "提交 5：Android 流式重组与设备版本化缓存。",
    "提交 6：Super-LIO operation 状态机、停止复核和定位质量门槛。",
    "提交 7：4KB/16KB A/B 与资源上限调优；决定是否需要 C++。",
])

doc.add_heading("7. 测试与验收方案", level=1)
doc.add_heading("7.1 测试分层", level=2)
add_table(
    doc,
    ["层级", "设备端", "Android", "关键产物"],
    [
        ["单元", "帧解析/CRC、队列策略、缓存key、operation状态机", "parser、single-writer、chunk assembler、超时/上限", "JUnit/Python测试报告"],
        ["协议契约", "Python/C++生成帧", "Kotlin解析并反向生成", "golden corpus + 字节级对比"],
        ["集成", "scheduler+模拟socket+ROS stub", "真机/模拟器连接设备模拟器", "端到端时序和错误码"],
        ["系统", "RK3588+MID-360S+ROS节点", "目标安卓设备", "现场地图/定位/导航记录"],
        ["故障注入", "节点挂死、磁盘满、服务超时、网络抖动", "断网、切后台、低内存、重复请求", "恢复与终态证明"],
        ["稳定性", "24h建图/定位/地图传输循环", "24h前后台/重连/下载", "RSS、线程、FD、队列趋势"],
    ],
    widths=[0.85, 2.65, 2.55, 1.25],
    font_size=7.8,
)

doc.add_heading("7.2 必测用例矩阵", level=2)
add_table(
    doc,
    ["编号", "场景", "注入/规模", "期望结果"],
    [
        ["T01", "TCP 分片与粘包", "每1字节、随机分片、100帧一次read", "帧数/顺序/载荷完全一致"],
        ["T02", "坏帧恢复", "错CRC、错尾、随机噪声、截断", "丢弃坏帧并在下一个STX恢复；无死循环/OOM"],
        ["T03", "并发发送", "心跳+急停+状态请求+地图", "线上字节流帧边界完整；急停不被大图阻塞"],
        ["T04", "地图乱序/重复/缺片", "2、10、50MB；随机重排、重复、丢1%", "重复幂等；缺片超时；不发布半图"],
        ["T05", "并发地图请求", "同map不同version、同version两个transfer", "按transfer_id隔离，无串图"],
        ["T06", "地图完整性", "篡改1字节、错误total_bytes/hash", "VERIFY_FAILED，临时文件清理"],
        ["T07", "连接中断", "传输25%/90%断网并重连", "旧transfer终止；新请求可成功；无泄漏"],
        ["T08", "生命周期启动失败", "缺服务、TF无效、话题旧时间戳", "终态FAILED/TIMEOUT，不报告READY"],
        ["T09", "生命周期停止失败", "子节点忽略shutdown", "仅kill自有节点；复核失败则ERROR并列残留"],
        ["T10", "定位质量", "空地图、低fitness、协方差大、位姿跳变", "拒绝READY，APP显示明确原因"],
        ["T11", "资源攻击/异常", "伪造超大total_chunks/total_bytes", "立即拒绝，RSS和磁盘受限"],
        ["T12", "24h稳定性", "周期保存/定位/下载/重连", "RSS、FD、线程和队列无持续增长"],
    ],
    widths=[0.48, 1.45, 2.65, 2.62],
    font_size=7.5,
)

doc.add_heading("7.3 建议验收指标（基线后确认）", level=2)
add_table(
    doc,
    ["指标", "建议目标", "测量口径"],
    [
        ["控制延迟", "局域网并发传图时 p95 <100ms，p99 <200ms；急停应用层p99 <100ms", "APP提交到设备应用确认；另测底盘实际停机时间"],
        ["地图吞吐", "相对当前基线≥3×；稳定网络下有效吞吐≥5MB/s", "不含首次地图生成；同时单列首包延迟"],
        ["首包延迟", "缓存命中 p95 <200ms；未命中单列PNG生成耗时", "设备收到请求到首个chunk写出"],
        ["正确性", "10万帧无交错、无错序；1000次地图传输摘要100%一致", "双端记录seq/transfer/hash"],
        ["资源", "单次传输额外RSS <64MB或<1.5×对象大小；队列有硬上限", "RK3588 RSS + Android PSS/GC"],
        ["生命周期", "启动/停止/保存/定位无假成功；所有请求有唯一终态", "故障注入100次"],
        ["稳定性", "24h无崩溃；RSS、线程、FD呈平台而非单调增长", "每分钟采样并做趋势图"],
    ],
    widths=[1.2, 2.5, 3.5],
    font_size=7.8,
)
p = doc.add_paragraph()
r = p.add_run("说明：")
r.bold = True
r.font.color.rgb = RED
p.add_run("以上是建议工程目标，不是当前实测结论。第一次基线完成后，应根据目标硬件、Wi‑Fi环境、地图尺寸和安全要求冻结正式SLO。")

doc.add_heading("7.4 性能测试组合", level=2)
add_table(
    doc,
    ["变量", "取值"],
    [
        ["地图对象", "2MB、10MB、50MB；PNG高/低压缩率；栅格原始格式"],
        ["chunk", "4KB、16KB；32KB仅探索，不直接生产默认"],
        ["网络", "无损；20/50/100ms RTT；1%丢包；5/20/100Mbps限速；突发抖动"],
        ["并发", "仅地图；地图+状态；地图+手动控制；地图+急停；双地图请求"],
        ["日志", "生产级；诊断级（用于量化日志开销）"],
        ["设备模式", "建图、定位、导航；回环开/关需分别记录"],
    ],
    widths=[1.4, 5.8],
)

doc.add_heading("8. 安全、兼容与回滚", level=1)
add_bullets(doc, [
    "控制面必须鉴权；生产凭据从仓库迁出到安全配置/密钥存储，外部链路采用TLS或受控网络隔离。",
    "协议新增字段只做向后兼容扩展；旧端不认识字段时可忽略，新端检测旧能力后回退4KB和旧组包逻辑。",
    "发送器、流式地图、C++ codec 均由独立 feature flag 控制，支持现场一键回退。",
    "急停、速度清零、task_enable=false 和底盘反馈不受大数据限流影响；故障时默认安全。",
    "灰度期间按设备维度保存版本、构建哈希、协议能力和关键指标；异常自动停止扩量。",
])

doc.add_heading("9. 交付物与上线门禁", level=1)
add_table(
    doc,
    ["交付物", "内容", "完成定义"],
    [
        ["基线报告", "CPU火焰图、Perfetto、端到端延迟、吞吐、RSS/GC", "场景可复现、原始数据可追溯"],
        ["协议版本", "proto、Python/Kotlin/C++生成物、能力协商、golden corpus", "三端字节兼容测试通过"],
        ["设备端改造", "bulk地图、流式窗口、缓存、operation状态机、健康门槛", "Python单测+ROS集成+RK3588系统测试"],
        ["Android改造", "单写者、流式组包、校验/取消/超时、日志治理", "JUnit+instrumented+真机弱网测试"],
        ["运维能力", "指标、诊断开关、残留节点清单、版本/feature flag", "现场可定位、可回滚"],
        ["验收包", "矩阵结果、24h报告、缺陷清单、上线与回滚手册", "P0清零，P1有明确豁免"],
    ],
    widths=[1.25, 3.35, 2.6],
    font_size=7.8,
)

add_callout(
    doc,
    "上线门禁",
    "未完成 Android 单写者、地图资源上限、生命周期无假成功、急停并发传图验证和24小时稳定性测试，不建议以“性能优化版”进入正式现场。",
    fill="FCE4D6",
)

doc.add_heading("附录 A：关键代码位置与本次核对结论", level=1)
add_table(
    doc,
    ["仓库", "文件", "关注点"],
    [
        ["设备端", "catkin_ws/src/grinder_scheduler/src/grinder_scheduler/sl_linka_adapter.py", "32KB接收、队列、send_lock、bulk发送；地图尚未接入bulk"],
        ["设备端", "catkin_ws/src/grinder_scheduler/src/grinder_scheduler/scheduler_node.py", "build_map_chunks：4KB上限、重复图像处理、完整outputs列表"],
        ["设备端", "catkin_ws/src/grinder_scheduler/src/grinder_scheduler/super_lio_mode_manager.py", "串行生命周期、停止兜底、READY门槛"],
        ["设备端", "third_party/sl_linka/proto/sl_link.proto", "MapChunk缺transfer_id/total_bytes/hash"],
        ["设备端", "catkin_ws/src/mapping/cloud_to_occupancy_grid/launch/mid360_mapping.launch", "当前包含Super-LIO/mapper/loop，不直接启动Livox驱动"],
        ["Android", "core/tcp/.../TcpService.kt", "多协程直接写OutputStream；4096B接收缓冲"],
        ["Android", "core/tcp/.../TcpManager.kt", "全帧hex日志；MapRequest请求4KB"],
        ["Android", "core/tcp/.../SlLinkManager.kt", "MapChunk按map_id全内存重组，无资源/超时/校验边界"],
        ["Android", "core/sllink/.../SlFrameParser.kt", "逐字节解析正确但CRC/载荷存在额外复制"],
    ],
    widths=[0.8, 3.8, 2.6],
    font_size=7.3,
)

doc.add_heading("附录 B：最终建议", level=1)
p = doc.add_paragraph()
p.add_run("建议批准路线一，暂缓路线二立项。").bold = True
p.add_run(" 先用 4–5 周把 Android 单写者、地图流式传输、协议事务字段、生命周期健康门槛和基准体系做好。这些改动同时提升速度、正确性和可维护性。完成后若设备端帧处理仍是明确热点，再用一个边界清晰、可回退的 C++ codec 扩展解决；不要先改写整个 scheduler，也不要先给 Android 引入 JNI。")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("— 文档结束 —")
r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)

doc.core_properties.title = "Grinder设备端与安卓端Python保持路线与C++热路径加速实施方案"
doc.core_properties.subject = "通信性能、地图传输、控制响应与ROS生命周期工程方案"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "Grinder, Android, ROS1, SL-LinkA, Super-LIO, 性能优化, C++"
doc.save(OUT)
print(OUT)
